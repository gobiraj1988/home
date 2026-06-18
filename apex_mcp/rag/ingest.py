"""Document loading & chunking for the APEX RAG index.

Responsibilities:
    * ``load_text``      -- read a single file to plain text, dispatching by extension.
    * ``iter_documents`` -- walk a directory tree yielding ``(path, text)`` pairs.
    * ``chunk_text``     -- split a document's text into overlapping chunks.

Every per-file operation is wrapped so that one unreadable/corrupt document can
never abort a full ingest -- failures are logged and skipped.
"""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Iterator

from ..common import get_logger

logger = get_logger(__name__)

# Directory name used by the persistent index; never ingest it.
INDEX_DIR_NAME = ".apex_index"


# --------------------------------------------------------------------------- #
# Per-extension loaders
# --------------------------------------------------------------------------- #


def _load_plain(path: Path) -> str:
    """Read a UTF-8 text file, tolerating decode errors."""
    return path.read_text(encoding="utf-8", errors="replace")


def _load_pdf(path: Path) -> str:
    """Extract text from a PDF using pypdf (imported lazily)."""
    from pypdf import PdfReader  # local import: keeps module import cheap

    reader = PdfReader(str(path))
    parts: list[str] = []
    for page_no, page in enumerate(reader.pages):
        try:
            parts.append(page.extract_text() or "")
        except Exception as exc:  # one bad page shouldn't kill the doc
            logger.warning("PDF page %d of %s failed: %s", page_no, path, exc)
    return "\n".join(parts)


def _load_docx(path: Path) -> str:
    """Extract text from a Word .docx using python-docx (imported lazily)."""
    import docx  # python-docx exposes the ``docx`` module

    document = docx.Document(str(path))
    lines = [para.text for para in document.paragraphs]
    # Include table cell text, which is otherwise lost.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _load_csv(path: Path, max_head: int = 10) -> str:
    """Summarise a CSV as readable text: shape, columns, dtypes, head, describe.

    We return a *summary* rather than the raw bytes so that wide/long tables stay
    retrievable without flooding the index with millions of cell values.
    """
    import pandas as pd  # imported lazily

    df = pd.read_csv(path)
    lines: list[str] = [
        f"CSV file: {path.name}",
        f"Shape: {df.shape[0]} rows x {df.shape[1]} columns",
        "",
        "Columns and dtypes:",
    ]
    for col, dtype in df.dtypes.items():
        lines.append(f"  - {col}: {dtype}")

    lines.append("")
    lines.append(f"Head ({min(max_head, len(df))} rows):")
    lines.append(df.head(max_head).to_string(index=False))

    # describe() only makes sense if there is at least one numeric column.
    numeric = df.select_dtypes(include="number")
    if not numeric.empty:
        lines.append("")
        lines.append("Numeric summary (describe):")
        lines.append(numeric.describe().to_string())

    return "\n".join(lines)


# Map of lowercase extension -> loader callable.
_LOADERS = {
    ".txt": _load_plain,
    ".md": _load_plain,
    ".markdown": _load_plain,
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".csv": _load_csv,
}


def load_text(path: Path) -> str:
    """Load a document to plain text, dispatching on file extension.

    Raises a ``ValueError`` for unsupported extensions; any parsing error from
    the underlying library propagates to the caller (callers in ``iter_documents``
    catch and skip). Returns the extracted text (possibly empty).
    """
    ext = path.suffix.lower()
    loader = _LOADERS.get(ext)
    if loader is None:
        raise ValueError(f"Unsupported file extension for load_text: {ext!r}")
    return loader(path)


# --------------------------------------------------------------------------- #
# Tree walking
# --------------------------------------------------------------------------- #


def _is_hidden(part: str) -> bool:
    """True for dotted directory/file names (e.g. ``.git``), excluding ``.``/``..``."""
    return part.startswith(".") and part not in (".", "..")


def iter_documents(root: Path, extensions: list[str]) -> Iterator[tuple[Path, str]]:
    """Walk ``root`` yielding ``(path, text)`` for every supported document.

    Skips the persistent index directory and any hidden directory/file. Per-file
    parsing errors are logged and skipped so a single bad document never aborts
    a full ingest.
    """
    # Normalise extensions to lowercase with a leading dot for comparison.
    wanted = {e.lower() if e.startswith(".") else f".{e.lower()}" for e in extensions}
    root = root.resolve()

    if not root.exists():
        logger.warning("RAG root does not exist: %s", root)
        return

    # os.walk-style traversal via rglob would not let us prune hidden dirs early,
    # so we walk manually with a stack to skip whole subtrees cheaply.
    stack: list[Path] = [root]
    while stack:
        current = stack.pop()
        try:
            entries = list(current.iterdir())
        except (PermissionError, OSError) as exc:
            logger.warning("Cannot list %s: %s", current, exc)
            continue

        for entry in entries:
            name = entry.name
            if entry.is_dir():
                if name == INDEX_DIR_NAME or _is_hidden(name):
                    continue  # prune index + hidden subtrees
                stack.append(entry)
                continue

            # It's a file.
            if _is_hidden(name):
                continue
            if entry.suffix.lower() not in wanted:
                continue

            try:
                text = load_text(entry)
            except Exception as exc:  # robust: log + skip any unreadable doc
                logger.warning("Skipping unreadable document %s: %s", entry, exc)
                continue

            if text and text.strip():
                yield entry, text
            else:
                logger.debug("Skipping empty document %s", entry)


# --------------------------------------------------------------------------- #
# Chunking
# --------------------------------------------------------------------------- #


def _chunk_id(source: str, chunk_index: int, text: str) -> str:
    """Stable, content-addressed id so re-ingesting unchanged docs is idempotent."""
    digest = hashlib.sha1(f"{source}:{chunk_index}:{text}".encode("utf-8")).hexdigest()
    return digest[:16]


def chunk_text(
    text: str,
    source: str,
    chunk_size: int = 1200,
    overlap: int = 200,
) -> list[dict]:
    """Split ``text`` into overlapping character windows.

    Args:
        text:       The full document text.
        source:     Relative path of the source doc (stored on each chunk).
        chunk_size: Target window size in characters.
        overlap:    Number of characters shared between consecutive windows so
                    that context spanning a boundary is still retrievable.

    Returns:
        A list of dicts with keys ``id``, ``source``, ``chunk_index``, ``text``.
    """
    if chunk_size <= 0:
        raise ValueError("chunk_size must be positive")
    # Clamp overlap so the window always advances (avoids an infinite loop).
    overlap = max(0, min(overlap, chunk_size - 1))

    text = text.strip()
    if not text:
        return []

    chunks: list[dict] = []
    step = chunk_size - overlap
    start = 0
    chunk_index = 0
    length = len(text)

    while start < length:
        end = min(start + chunk_size, length)
        piece = text[start:end].strip()
        if piece:
            chunks.append(
                {
                    "id": _chunk_id(source, chunk_index, piece),
                    "source": source,
                    "chunk_index": chunk_index,
                    "text": piece,
                }
            )
            chunk_index += 1
        if end >= length:
            break
        start += step

    return chunks
